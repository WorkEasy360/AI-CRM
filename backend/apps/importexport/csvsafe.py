"""CSV/XLSX safety: upload validation, size/row caps and spreadsheet formula neutralisation (T12, T13, T19).

An XLSX upload is converted to CSV once, at upload, and every later step (preview, mapping, import,
resume) reads that CSV, so the rest of the pipeline knows one format.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile
from collections.abc import Iterator
from decimal import Decimal, InvalidOperation
from typing import Any

from rest_framework.exceptions import ValidationError

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 60
MAX_HEADER_LENGTH = 100
MAX_CELL_LENGTH = 5_000
SNIFF_BYTES = 8 * 1024

# A 5 MB .xlsx can inflate a lot: refuse workbooks whose XML expands beyond this (zip-bomb guard).
MAX_XLSX_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r", "\n")
_NUMBER_RE = re.compile(r"^-?\d+(\.\d+)?$")
_FILENAME_SAFE = re.compile(r"[^A-Za-z0-9._ -]")

csv.field_size_limit(MAX_CELL_LENGTH * 4)


def neutralise(value: Any) -> str:
    """Make a cell inert for spreadsheet applications.

    Cells that a spreadsheet would evaluate (leading ``= + - @`` or control characters) are prefixed
    with a single quote; plain negative numbers are left alone so numeric columns stay usable. Pipes
    are neutralised as well because some applications treat ``|`` as a DDE separator.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, Decimal | int | float):
        return str(value)
    text = str(value)
    if not text:
        return ""
    if text[0] in FORMULA_TRIGGERS and not _NUMBER_RE.match(text):
        text = "'" + text
    if "|" in text and text[0] != "'":
        text = "'" + text
    return text.replace("\x00", "")


def safe_filename(name: str, *, default: str = "upload.csv") -> str:
    name = (name or "").replace("\\", "/").rsplit("/", 1)[-1]
    name = _FILENAME_SAFE.sub("_", name).strip(" .")
    return (name or default)[:120]


def validate_upload(uploaded: Any) -> tuple[bytes, list[str], int]:
    """Validate an uploaded CSV. Returns (raw bytes, headers, data row count).

    Checks: extension, size, binary content (NUL bytes / non-UTF-8), header shape, column and row caps,
    cell length cap. The whole file is parsed once so a malformed file is rejected before it is stored.
    """
    name = safe_filename(getattr(uploaded, "name", ""))
    is_xlsx = name.lower().endswith(".xlsx")
    if not (name.lower().endswith(".csv") or is_xlsx):
        raise ValidationError({"file": "Upload a .csv or .xlsx file."})
    size = getattr(uploaded, "size", None)
    if size is None or size <= 0:
        raise ValidationError({"file": "The file is empty."})
    if size > MAX_UPLOAD_BYTES:
        raise ValidationError({"file": f"The file is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB."})
    raw = uploaded.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValidationError({"file": "The file is too large."})
    if is_xlsx:
        raw = xlsx_to_csv(raw)
    head = raw[:SNIFF_BYTES]
    if b"\x00" in head:
        raise ValidationError({"file": "The file does not look like a text CSV file."})
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValidationError({"file": "The file must be UTF-8 encoded."}) from exc
    headers, count = _scan(text)
    return raw, headers, count


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.isoformat() if value.time() != dt.time(0) else value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def xlsx_to_csv(raw: bytes) -> bytes:
    """The first worksheet of an .xlsx workbook as UTF-8 CSV (values only, never formulas)."""
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValidationError({"file": "The workbook is too large once unpacked."})
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except ValidationError:
        raise
    except Exception as exc:  # openpyxl raises a zoo of exception types for malformed files
        raise ValidationError({"file": "The file is not a valid .xlsx workbook."}) from exc
    try:
        sheet = workbook.worksheets[0]
        out = io.StringIO(newline="")
        writer = csv.writer(out, lineterminator="\r\n")
        rows = 0
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell_text(v) for v in row[: MAX_COLUMNS + 1]]
            while cells and cells[-1] == "":
                cells.pop()
            if not cells:
                continue
            rows += 1
            if rows > MAX_ROWS + 1:
                raise ValidationError({"file": f"At most {MAX_ROWS} rows per import. Split the file."})
            writer.writerow(cells)
    finally:
        workbook.close()
    return out.getvalue().encode("utf-8")


def _scan(text: str) -> tuple[list[str], int]:
    reader = csv.reader(io.StringIO(text, newline=""))
    try:
        headers = next(reader)
    except StopIteration as exc:
        raise ValidationError({"file": "The file has no header row."}) from exc
    except csv.Error as exc:
        raise ValidationError({"file": "The file is not valid CSV."}) from exc
    headers = [h.strip().replace("\x00", "") for h in headers]
    if not headers or all(not h for h in headers):
        raise ValidationError({"file": "The header row is empty."})
    if len(headers) > MAX_COLUMNS:
        raise ValidationError({"file": f"At most {MAX_COLUMNS} columns are supported."})
    if any(len(h) > MAX_HEADER_LENGTH for h in headers):
        raise ValidationError({"file": "A column header is too long."})
    if len({h.lower() for h in headers if h}) != len([h for h in headers if h]):
        raise ValidationError({"file": "Column headers must be unique."})
    count = 0
    try:
        for row in reader:
            if not any(cell.strip() for cell in row):
                continue
            count += 1
            if count > MAX_ROWS:
                raise ValidationError({"file": f"At most {MAX_ROWS} rows per import. Split the file."})
            if len(row) > len(headers):
                raise ValidationError({"file": f"Row {count + 1} has more cells than the header."})
            if any(len(cell) > MAX_CELL_LENGTH for cell in row):
                raise ValidationError({"file": f"Row {count + 1} has a cell longer than {MAX_CELL_LENGTH} characters."})
    except csv.Error as exc:
        raise ValidationError({"file": "The file is not valid CSV."}) from exc
    if count == 0:
        raise ValidationError({"file": "The file has no data rows."})
    return headers, count


def iter_rows(raw: bytes, headers: list[str]) -> Iterator[tuple[int, dict[str, str]]]:
    """Yield (1-based data row number, {header: cell}) skipping blank rows."""
    reader = csv.reader(io.StringIO(raw.decode("utf-8-sig"), newline=""))
    next(reader, None)
    n = 0
    for row in reader:
        if not any(cell.strip() for cell in row):
            continue
        n += 1
        yield n, {headers[i]: (row[i] if i < len(row) else "") for i in range(len(headers)) if headers[i]}


def parse_decimal(value: str) -> Decimal | None:
    value = (value or "").strip().replace(",", "")
    if not value:
        return None
    try:
        return Decimal(value)
    except InvalidOperation as exc:
        raise ValidationError("Expected a number.") from exc
